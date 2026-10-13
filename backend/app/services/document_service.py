"""Document processing service — works with both Celery and FastAPI BackgroundTasks.

In lite mode (LITE_MODE=true), document processing runs inline via FastAPI's
BackgroundTasks. In full mode, it runs via Celery workers.

This module handles the full pipeline:
1. Text extraction (PDF, DOCX, TXT, XLSX, CSV)
2. Chunking with overlap
3. Embedding generation and ChromaDB storage
"""

from __future__ import annotations

import hashlib
import time
from typing import Any

from app.core.constants import DEFAULT_CHUNK_OVERLAP, DEFAULT_CHUNK_SIZE
from app.core.logging import get_logger

logger = get_logger(__name__)


async def process_document_async(
    document_id: str,
    tenant_id: str,
    file_path: str,
    file_type: str,
) -> dict[str, Any]:
    """Process an uploaded document asynchronously (lite mode).

    This is the async version used with FastAPI BackgroundTasks
    when Celery is not available.

    Args:
        document_id: UUID of the Document record.
        tenant_id: Tenant UUID for scoping.
        file_path: Path to the uploaded file.
        file_type: Document type (pdf, docx, txt, etc.).

    Returns:
        Processing result with chunk count and metadata.
    """
    start_time = time.time()

    try:
        logger.info(
            "Processing document (async)",
            document_id=document_id,
            file_type=file_type,
        )

        # Step 1: Extract text
        text = _extract_text(file_path, file_type)
        if not text:
            await _update_document_status(document_id, "failed", error="No text could be extracted")
            return {
                "status": "failed",
                "document_id": document_id,
                "error": "No text could be extracted",
            }

        # Step 2: Compute content hash for deduplication
        content_hash = hashlib.sha256(text.encode()).hexdigest()

        # Step 3: Chunk the text
        chunks = _chunk_text(text)

        # Step 4: Store chunks in ChromaDB
        stored_count = await _store_chunks_in_chromadb(
            chunks=chunks,
            tenant_id=tenant_id,
            document_id=document_id,
            file_path=file_path,
            file_type=file_type,
        )

        duration_ms = int((time.time() - start_time) * 1000)

        # Step 5: Update document record
        await _update_document_status(
            document_id,
            status="completed",
            chunk_count=stored_count,
            character_count=len(text),
            content_hash=content_hash,
        )

        result = {
            "status": "completed",
            "document_id": document_id,
            "content_hash": content_hash,
            "character_count": len(text),
            "chunk_count": stored_count,
            "duration_ms": duration_ms,
        }

        logger.info(
            "Document processed successfully",
            document_id=document_id,
            chunk_count=stored_count,
            duration_ms=duration_ms,
        )

        return result

    except Exception as exc:
        logger.error(
            "Document processing failed",
            document_id=document_id,
            error=str(exc),
            exc_info=True,
        )
        await _update_document_status(document_id, "failed", error=str(exc))
        return {
            "status": "failed",
            "document_id": document_id,
            "error": str(exc),
        }


async def _store_chunks_in_chromadb(
    chunks: list[str],
    tenant_id: str,
    document_id: str,
    file_path: str,
    file_type: str,
) -> int:
    """Store text chunks in ChromaDB with metadata.

    Each chunk is stored with metadata including source filename,
    page number, and chunk index for citation tracking.

    Args:
        chunks: List of text chunks to store.
        tenant_id: Tenant ID for collection scoping.
        document_id: Source document UUID.
        file_path: Original file path (used for source metadata).
        file_type: Document type.

    Returns:
        Number of chunks successfully stored.
    """
    if not chunks:
        return 0

    try:
        import chromadb

        from app.core.config import get_settings

        settings = get_settings()
        client = chromadb.HttpClient(
            host=settings.chromadb_host,
            port=settings.chromadb_port,
        )

        # Create or get the tenant-scoped collection
        collection_name = f"tenant_{tenant_id}_documents"
        collection = client.get_or_create_collection(
            name=collection_name,
            metadata={"hnsw:space": "cosine"},
        )

        # Extract source filename from path
        from pathlib import Path
        source_name = Path(file_path).name

        # Prepare data for batch upsert
        ids = []
        documents = []
        metadatas = []

        for i, chunk in enumerate(chunks):
            chunk_id = f"{document_id}_chunk_{i}"
            ids.append(chunk_id)
            documents.append(chunk)
            metadatas.append({
                "source": source_name,
                "document_id": document_id,
                "chunk_index": i,
                "chunk_id": chunk_id,
                "file_type": file_type,
                "tenant_id": tenant_id,
                "page_number": _estimate_page_number(i, len(chunks), file_type),
            })

        # ChromaDB handles embedding internally using its default model
        # (all-MiniLM-L6-v2 via sentence-transformers)
        collection.upsert(
            ids=ids,
            documents=documents,
            metadatas=metadatas,
        )

        logger.info(
            "Chunks stored in ChromaDB",
            collection=collection_name,
            chunk_count=len(chunks),
            document_id=document_id,
        )

        return len(chunks)

    except Exception as exc:
        logger.error(
            "ChromaDB storage failed",
            error=str(exc),
            document_id=document_id,
            exc_info=True,
        )
        return 0


async def _update_document_status(
    document_id: str,
    status: str,
    error: str = "",
    chunk_count: int = 0,
    character_count: int = 0,
    content_hash: str = "",
) -> None:
    """Update the document record in the database with processing results."""
    try:
        import uuid

        from sqlalchemy import update

        from app.database.session import get_session_factory
        from app.models.document import Document

        factory = get_session_factory()
        async with factory() as session:
            values: dict[str, Any] = {"status": status}
            if error:
                values["error_message"] = error
            if chunk_count:
                values["chunk_count"] = chunk_count
            if character_count:
                values["character_count"] = character_count
            if content_hash:
                values["content_hash"] = content_hash

            await session.execute(
                update(Document)
                .where(Document.id == uuid.UUID(document_id))
                .values(**values)
            )
            await session.commit()
    except Exception as exc:
        logger.error("Failed to update document status", error=str(exc), exc_info=True)


def _extract_text(file_path: str, file_type: str) -> str:
    """Extract text content from a document file.

    Args:
        file_path: Path to the file.
        file_type: Document type.

    Returns:
        Extracted text content.
    """
    ft = file_type.lower()
    if ft == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(file_path)
        return "\n".join(
            page.extract_text() or "" for page in reader.pages
        )
    elif ft == "docx":
        import docx

        doc = docx.Document(file_path)
        return "\n".join(para.text for para in doc.paragraphs if para.text)
    elif ft in ("txt", "csv"):
        with open(file_path, encoding="utf-8") as f:
            return f.read()
    elif ft == "xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(file_path, data_only=True)
        text_parts: list[str] = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join(str(cell) for cell in row if cell is not None)
                if row_text:
                    text_parts.append(row_text)
        return "\n".join(text_parts)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")


def _chunk_text(
    text: str,
    chunk_size: int = DEFAULT_CHUNK_SIZE,
    chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> list[str]:
    """Split text into overlapping chunks.

    Uses a simple character-based splitter with paragraph-aware
    boundary detection.

    Args:
        text: The full text to chunk.
        chunk_size: Target chunk size in characters.
        chunk_overlap: Overlap between consecutive chunks.

    Returns:
        List of text chunks.
    """
    if len(text) <= chunk_size:
        return [text]

    chunks: list[str] = []
    start = 0

    while start < len(text):
        end = start + chunk_size

        # Try to break at paragraph or sentence boundary
        if end < len(text):
            for sep in ["\n\n", "\n", ". ", "! ", "? "]:
                last_sep = text.rfind(sep, start, end)
                if last_sep > start + chunk_size // 2:
                    end = last_sep + len(sep)
                    break

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        start = end - chunk_overlap

    return chunks


def _estimate_page_number(chunk_index: int, total_chunks: int, file_type: str) -> int:
    """Estimate page number for a chunk (rough approximation).

    For PDFs, assumes ~3 chunks per page. For other formats, returns 1.
    """
    if file_type == "pdf":
        return (chunk_index // 3) + 1
    return 1
